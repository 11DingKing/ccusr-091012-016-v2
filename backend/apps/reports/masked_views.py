"""
脱敏报表视图

- 用途列表：按角色返回可选用途及其处理规则
- 趋势预览：JSON 形式发布脱敏结果（留痕）
- 趋势导出：xlsx 形式发布，附“导出摘要”页供审核人确认（留痕）
- 导出记录：审核人查看历次预览/导出的策略版本与数据范围
"""
import logging
from datetime import datetime, timedelta

from django.http import HttpResponse
from django.utils import timezone
from django.utils.http import content_disposition_header
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from rest_framework.permissions import IsAuthenticated
from rest_framework.views import APIView

from apps.core.response import error_response, success_response
from .masking import (
    ROW_KIND_LABELS,
    MaskedReportError,
    available_purposes,
    build_masked_report,
    record_export_audit,
)
from .models import MaskedExportRecord

logger = logging.getLogger('apps')


def _parse_date(value, field_name):
    try:
        return datetime.strptime(value, '%Y-%m-%d').date()
    except (TypeError, ValueError):
        raise MaskedReportError(f'{field_name} 格式应为 YYYY-MM-DD')


def _extract_common_params(request):
    """解析并校验报表公共参数，返回 (purpose, start, end, granularity, dimension)。"""
    params = request.query_params
    purpose = params.get('purpose')
    if not purpose:
        raise MaskedReportError('缺少参数：purpose（报表用途）')

    today = timezone.localdate()
    end = _parse_date(params.get('end_date'), 'end_date') if params.get('end_date') else today
    start = _parse_date(params.get('start_date'), 'start_date') if params.get('start_date') else end - timedelta(days=30)

    granularity = params.get('granularity') or None
    dimension = params.get('dimension') or None
    return purpose, start, end, granularity, dimension


def _build_for_request(request):
    purpose, start, end, granularity, dimension = _extract_common_params(request)
    report = build_masked_report(
        request.user, purpose, start, end,
        granularity=granularity, dimension=dimension,
    )
    return purpose, start, end, report


class MaskedReportPurposesView(APIView):
    """列出当前角色可用的脱敏报表用途及处理规则"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return success_response(data=available_purposes(request.user.role))


class MaskedTrendReportView(APIView):
    """脱敏收发趋势预览（JSON），每次预览均留痕"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        try:
            purpose, start, end, report = _build_for_request(request)
        except MaskedReportError as exc:
            return error_response(message=exc.message, code=exc.code)

        record_export_audit(request.user, 'preview', purpose, report, start, end)
        logger.info(f"Masked report previewed: {purpose} by {request.user.username}")
        return success_response(data={
            'summary': report['summary'],
            'rows': report['rows'],
            'totals': report['totals'],
        })


class MaskedTrendExportView(APIView):
    """脱敏收发趋势导出（xlsx），附导出摘要页，每次导出均留痕"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        try:
            purpose, start, end, report = _build_for_request(request)
        except MaskedReportError as exc:
            return error_response(message=exc.message, code=exc.code)

        record_export_audit(request.user, 'export', purpose, report, start, end)

        wb = Workbook()
        self._build_summary_sheet(wb.active, report['summary'])
        self._build_data_sheet(wb.create_sheet('收发趋势'), report)

        fingerprint = report['summary']['data_fingerprint']
        filename = f"脱敏收发趋势_{purpose}_{report['summary']['policy_version']}_{fingerprint[:8]}.xlsx"
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = content_disposition_header(True, filename)
        wb.save(response)

        logger.info(f"Masked report exported: {purpose} by {request.user.username}")
        return response

    def _build_summary_sheet(self, ws, summary):
        ws.title = '导出摘要'
        ws.append(['项目', '内容'])
        pairs = [
            ('报表名称', summary['report_title']),
            ('报表用途', f"{summary['purpose_label']}（{summary['purpose']}）"),
            ('用途说明', summary['purpose_description']),
            ('策略版本', summary['policy_version']),
            ('请求数据范围', f"{summary['requested_window']['start']} ~ {summary['requested_window']['end']}"),
            ('实际数据范围', f"{summary['effective_window']['start']} ~ {summary['effective_window']['end']}"),
            ('窗口调整说明', '；'.join(summary['window_notes'])),
            ('冻结规则', summary['freeze_rule']),
            ('统计口径', summary['measure_rule']),
            ('时间粒度', summary['granularity_label']),
            ('统计维度', summary['dimension_label']),
            ('小样本阈值', f"k={summary['k_threshold']}"),
            ('小样本处理策略', summary['small_cell_strategy_label']),
            ('发布字段', '、'.join(summary['fields_released'])),
            ('裁剪字段（不发布）', '、'.join(summary['fields_withheld'])),
            ('时间桶数', summary['bucket_count']),
            ('发布行数', summary['row_count']),
            ('其中合并行数', summary['merged_row_count']),
            ('被抑制单元数', summary['suppressed_cell_count']),
            ('抑制说明', summary['suppressed_note']),
            ('合计口径', summary['totals_rule']),
            ('一致性说明', summary['consistency_note']),
            ('数据指纹', summary['data_fingerprint']),
            ('导出人', f"{summary['generated_by']}（{summary['generated_by_role']}）"),
            ('导出时间', summary['generated_at']),
        ]
        for label, value in pairs:
            ws.append([label, value])
        self._style_sheet(ws, header_rows=1)
        ws.column_dimensions['A'].width = 22
        ws.column_dimensions['B'].width = 90

    def _build_data_sheet(self, ws, report):
        summary = report['summary']
        headers = summary['fields_released']
        ws.append(headers)

        with_goods_count = '涉及货物种类数' in headers
        for row in report['rows']:
            line = [
                row['bucket'],
                row['dimension_value'],
                ROW_KIND_LABELS[row['row_kind']],
                row['in_count'] if row['in_count'] is not None else '',
                row['in_quantity'] if row['in_quantity'] is not None else '',
                row['out_count'] if row['out_count'] is not None else '',
                row['out_quantity'] if row['out_quantity'] is not None else '',
                row['record_count'] if row['record_count'] is not None else '',
            ]
            if with_goods_count:
                line.append(row['goods_type_count'] if row['goods_type_count'] is not None else '')
            ws.append(line)

        # 合计行：仅统计已发布单元
        grand = report['totals']['grand']
        total_line = [
            '合计（仅已发布单元）', '', '',
            grand['in_count'], grand['in_quantity'],
            grand['out_count'], grand['out_quantity'],
            grand['record_count'],
        ]
        if with_goods_count:
            total_line.append('')
        ws.append(total_line)

        self._style_sheet(ws, header_rows=1)
        for idx, header in enumerate(headers, start=1):
            ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = max(14, len(str(header)) * 2 + 4)

    def _style_sheet(self, ws, header_rows=1):
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='0066FF', end_color='0066FF', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin'),
        )
        for row in ws.iter_rows(min_row=1, max_row=header_rows):
            for cell in row:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border


class MaskedExportRecordListView(APIView):
    """脱敏报表预览/导出审计记录（仅管理角色可核对）"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        if not request.user.is_admin:
            return error_response(message='仅管理角色可查看导出审计记录', code=403)

        records = MaskedExportRecord.objects.all()[:200]
        data = []
        for record in records:
            data.append({
                'id': record.id,
                'action': record.action,
                'action_label': record.get_action_display(),
                'requested_by': record.requested_by.username if record.requested_by else '',
                'requester_role': record.requester_role,
                'purpose': record.purpose,
                'purpose_label': record.purpose_label,
                'policy_version': record.policy_version,
                'granularity': record.granularity,
                'dimension': record.dimension,
                'requested_window': {
                    'start': record.requested_start.isoformat(),
                    'end': record.requested_end.isoformat(),
                },
                'effective_window': {
                    'start': record.effective_start.isoformat(),
                    'end': record.effective_end.isoformat(),
                },
                'k_threshold': record.k_threshold,
                'strategy': record.strategy,
                'fields_released': record.fields_released,
                'fields_withheld': record.fields_withheld,
                'bucket_count': record.bucket_count,
                'row_count': record.row_count,
                'released_row_count': record.released_row_count,
                'merged_row_count': record.merged_row_count,
                'suppressed_cell_count': record.suppressed_cell_count,
                'suppressed_record_count': record.suppressed_record_count,
                'data_fingerprint': record.data_fingerprint,
                'summary': record.summary,
                'created_at': record.created_at.isoformat(),
            })
        return success_response(data=data)
