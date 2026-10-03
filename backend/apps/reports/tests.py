from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient, APIRequestFactory, force_authenticate

from apps.authentication.backends import generate_token
from apps.authentication.models import User
from apps.warehouse.models import (
    Category, Goods, StockIn, StockOut, Unit, Variety, Warning,
)
from .anonymization import STRATEGY_VERSION
from .cron import check_stock_warning
from .models import AnonymizedBucketSnapshot, DailyReport
from .views import DashboardView


class ReportTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("report-user", "testpass123", role="admin")
        self.client = APIClient()
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {generate_token(self.user)}")
        unit = Unit.objects.create(name="件", created_by=self.user)
        category = Category.objects.create(name="监管器材", unit=unit, created_by=self.user)
        variety = Variety.objects.create(name="封存设备", category=category, created_by=self.user)
        self.goods = Goods.objects.create(
            variety=variety, name="封存终端", code="RPT-001", quantity=Decimal("2"), warning_threshold=Decimal("5")
        )

    def test_daily_report_model(self):
        report = DailyReport.objects.create(report_date=date(2026, 9, 30), in_count=2, in_total=4, out_count=1, out_total=1)
        self.assertEqual(report.in_count, 2)
        self.assertIn("2026-09-30", str(report))

    def test_dashboard_view_counts_warning_goods(self):
        request = APIRequestFactory().get("/api/dashboard/")
        force_authenticate(request, user=self.user)
        response = DashboardView.as_view()(request)
        response.render()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["data"]["warning_goods_count"], 1)

    def test_warning_job_is_idempotent_for_unread_warning(self):
        check_stock_warning()
        check_stock_warning()
        self.assertEqual(Warning.objects.filter(goods=self.goods, is_read=False).count(), 1)

    def test_daily_report_range(self):
        response = self.client.get("/api/daily-report/", {"start_date": "2026-09-29", "end_date": "2026-10-01"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.json()["data"]), 3)

    def test_invalid_export_type(self):
        response = self.client.get("/api/export/", {"type": "unknown"})
        self.assertEqual(response.status_code, 400)

    def test_system_monitor_shape(self):
        response = self.client.get("/api/system-monitor/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("cpu", response.json()["data"])


class AnonymizedTrendReportTest(TestCase):
    """脱敏趋势报表：字段裁剪、小样本处理与防差分固化。"""

    def setUp(self):
        self.admin = User.objects.create_user("anon-admin", "testpass123", role="admin")
        self.outsider = User.objects.create_user("anon-user", "testpass123", role="user")
        self.client = APIClient()
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {generate_token(self.admin)}")

        unit = Unit.objects.create(name="件", created_by=self.admin)
        # 三个品类：A 组样本充足，B/C 组为小组样本
        self.cat_a = Category.objects.create(name="防护装备", unit=unit, created_by=self.admin)
        self.cat_b = Category.objects.create(name="通讯器材", unit=unit, created_by=self.admin)
        self.cat_c = Category.objects.create(name="检验耗材", unit=unit, created_by=self.admin)
        var_a = Variety.objects.create(name="防护服", category=self.cat_a, created_by=self.admin)
        var_b = Variety.objects.create(name="对讲机", category=self.cat_b, created_by=self.admin)
        var_c = Variety.objects.create(name="试剂盒", category=self.cat_c, created_by=self.admin)
        self.goods_a = Goods.objects.create(variety=var_a, name="连体防护服", code="ANON-A1", quantity=Decimal("100"))
        self.goods_b = Goods.objects.create(variety=var_b, name="集群对讲机", code="ANON-B1", quantity=Decimal("10"))
        self.goods_c = Goods.objects.create(variety=var_c, name="快检试剂盒", code="ANON-C1", quantity=Decimal("10"))

        # 以上一个完整自然周为数据周期，保证周期已结束
        today = timezone.now().date()
        self.week_start = today - timedelta(days=today.weekday() + 7)
        self.week_end = self.week_start + timedelta(days=6)

        # A 组：3 条入库（达到小样本阈值），含案件/人员敏感字段
        for i in range(3):
            self._make_stock_in(self.goods_a, Decimal("2.5"), self.week_start + timedelta(days=i),
                                batch_no=f"CASE-2026-{i}", supplier="某涉案供应商")
        # B 组：2 条入库（小组）
        for i in range(2):
            self._make_stock_in(self.goods_b, Decimal("1"), self.week_start + timedelta(days=i))
        # C 组：1 条出库（小组）
        self._make_stock_out(self.goods_c, Decimal("1"), self.week_start + timedelta(days=2),
                             receiver="张三", receiver_dept="专案一组")

    def _make_stock_in(self, goods, quantity, day, batch_no="", supplier=""):
        record = StockIn.objects.create(
            goods=goods, operator=self.admin, quantity=quantity,
            batch_no=batch_no, supplier=supplier, remark="涉案件备注"
        )
        aware = timezone.make_aware(datetime.combine(day, time(10, 0)))
        StockIn.objects.filter(pk=record.pk).update(stock_in_time=aware)
        return record

    def _make_stock_out(self, goods, quantity, day, receiver="", receiver_dept=""):
        return StockOut.objects.create(
            goods=goods, operator=self.admin, quantity=quantity, status="completed",
            receiver=receiver, receiver_dept=receiver_dept,
            stock_out_time=timezone.make_aware(datetime.combine(day, time(15, 0))),
            remark="涉案件备注"
        )

    def _get_report(self, user=None, **params):
        if user is not None:
            self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {generate_token(user)}")
        response = self.client.get("/api/anonymized-reports/trend/", params)
        self.assertEqual(response.status_code, 200, response.content)
        return response.json()["data"]

    def _week_params(self, **extra):
        params = {
            "purpose": "external_inspection",
            "granularity": "week",
            "start_date": self.week_start.isoformat(),
            "end_date": self.week_end.isoformat(),
        }
        params.update(extra)
        return params

    # ---------- 字段裁剪 ----------

    def test_external_purpose_strips_sensitive_fields(self):
        data = self._get_report(**self._week_params())
        self.assertTrue(data["rows"])
        allowed = {"period", "category", "in_count", "in_total", "out_count", "out_total"}
        for row in data["rows"]:
            self.assertLessEqual(set(row.keys()), allowed)
        body = str(data["rows"])
        for leaked in ("张三", "专案一组", "CASE-2026", "某涉案供应商", "涉案件备注",
                       "连体防护服", "ANON-A1", "anon-admin"):
            self.assertNotIn(leaked, body)

    def test_internal_audit_adds_variety_dimension(self):
        data = self._get_report(**self._week_params(purpose="internal_audit"))
        self.assertIn("variety", data["summary"]["dimensions"])
        self.assertTrue(any(row.get("variety") == "防护服" for row in data["rows"]))

    def test_internal_audit_rejects_plain_user(self):
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {generate_token(self.outsider)}")
        response = self.client.get("/api/anonymized-reports/trend/",
                                   self._week_params(purpose="internal_audit"))
        self.assertEqual(response.status_code, 403)

    # ---------- 小样本合并与抑制 ----------

    def test_small_groups_merged_into_other_bucket(self):
        data = self._get_report(**self._week_params())
        labels = {row["category"] for row in data["rows"]}
        self.assertIn("防护装备", labels)          # 大组保留
        self.assertNotIn("通讯器材", labels)        # 小组不单独出现
        self.assertNotIn("检验耗材", labels)
        self.assertIn("其他（小样本合并）", labels)  # 合并桶出现
        merged = [r for r in data["rows"] if r["category"] == "其他（小样本合并）"][0]
        self.assertEqual(merged["in_count"], 2)     # B 组两条入库
        self.assertEqual(merged["out_count"], 1)    # C 组一条出库
        self.assertEqual(data["summary"]["totals"]["merged_groups"], 2)

    def test_merged_bucket_suppressed_when_still_below_threshold(self):
        # 删除 A 组与 B 组数据，只剩 C 组 1 条出库：合并后仍不足阈值，整桶抑制
        StockIn.objects.all().delete()
        data = self._get_report(**self._week_params())
        self.assertEqual(data["rows"], [])
        self.assertEqual(data["summary"]["totals"]["suppressed_groups"], 1)

    def test_external_totals_are_rounded(self):
        data = self._get_report(**self._week_params())
        big = [r for r in data["rows"] if r["category"] == "防护装备"][0]
        # 3 × 2.5 = 7.5，按步长 1 确定性取整为 8
        self.assertEqual(big["in_total"], 8.0)
        self.assertTrue(float(big["in_total"]).is_integer())

    # ---------- 防差分：重叠窗口 / 迟到更正 / 跨用户 ----------

    def test_overlapping_windows_yield_identical_shared_periods(self):
        wide = self._get_report(**self._week_params(
            start_date=(self.week_start - timedelta(days=7)).isoformat()))
        narrow = self._get_report(**self._week_params())
        wide_rows = {(r["period"], r["category"]): r for r in wide["rows"]}
        for row in narrow["rows"]:
            key = (row["period"], row["category"])
            self.assertIn(key, wide_rows)
            self.assertEqual(wide_rows[key], row)

    def test_late_correction_does_not_change_published_snapshot(self):
        before = self._get_report(**self._week_params())
        # 迟到更正：向已发布周期补录一条记录
        self._make_stock_in(self.goods_a, Decimal("99"), self.week_start)
        after = self._get_report(**self._week_params())
        self.assertEqual(before["rows"], after["rows"])
        self.assertEqual(
            AnonymizedBucketSnapshot.objects.filter(
                purpose="external_inspection", granularity="week"
            ).count(), 1
        )

    def test_different_roles_get_identical_report(self):
        admin_view = self._get_report(user=self.admin, **self._week_params())
        user_view = self._get_report(user=self.outsider, **self._week_params())
        self.assertEqual(admin_view["rows"], user_view["rows"])
        strip = lambda s: {k: v for k, v in s.items() if k != "generated_at"}
        self.assertEqual(strip(admin_view["summary"]), strip(user_view["summary"]))

    # ---------- 窗口对齐与进行中周期 ----------

    def test_window_snaps_outward_to_period_bounds(self):
        data = self._get_report(**self._week_params(
            start_date=(self.week_start + timedelta(days=2)).isoformat(),
            end_date=(self.week_start + timedelta(days=4)).isoformat()))
        self.assertEqual(data["summary"]["aligned_window"]["start"], self.week_start.isoformat())
        self.assertEqual(data["summary"]["aligned_window"]["end"], self.week_end.isoformat())

    def test_current_period_is_excluded(self):
        today = timezone.now().date()
        data = self._get_report(**self._week_params(end_date=today.isoformat()))
        self.assertTrue(data["summary"]["current_period_excluded"])
        for period in data["summary"]["periods"]:
            self.assertLess(date.fromisoformat(period["end"]), today)

    # ---------- 导出摘要与审计 ----------

    def test_summary_discloses_scope_and_strategy(self):
        data = self._get_report(**self._week_params())
        summary = data["summary"]
        self.assertEqual(summary["strategy_version"], STRATEGY_VERSION)
        self.assertEqual(summary["strategy"]["min_group_size"], 3)
        self.assertEqual(summary["requested_window"]["start"], self.week_start.isoformat())
        self.assertIn("period", summary["fields"])
        self.assertTrue(summary["excluded_sensitive_fields"])
        self.assertTrue(summary["snapshot_ids"])
        self.assertEqual(len(summary["periods"]), 1)
        self.assertIn("data_frozen_at", summary["periods"][0])

    def test_policy_endpoint_lists_purposes_and_rules(self):
        response = self.client.get("/api/anonymized-reports/policy/")
        self.assertEqual(response.status_code, 200)
        data = response.json()["data"]
        self.assertEqual(data["strategy_version"], STRATEGY_VERSION)
        purposes = {p["purpose"]: p for p in data["purposes"]}
        self.assertIn("external_inspection", purposes)
        self.assertTrue(purposes["external_inspection"]["allowed"])

    def test_xlsx_export_contains_summary_sheet(self):
        response = self.client.get("/api/anonymized-reports/trend/",
                                   self._week_params(export="xlsx"))
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(BytesIO(response.content))
        self.assertIn("导出摘要", wb.sheetnames)
        self.assertIn("趋势数据", wb.sheetnames)
        summary_text = " ".join(
            str(cell.value) for row in wb["导出摘要"].iter_rows() for cell in row
        )
        self.assertIn(STRATEGY_VERSION, summary_text)
        self.assertIn("脱敏策略版本", summary_text)

    def test_invalid_purpose_and_granularity_rejected(self):
        response = self.client.get("/api/anonymized-reports/trend/",
                                   self._week_params(purpose="spy"))
        self.assertEqual(response.status_code, 400)
        response = self.client.get("/api/anonymized-reports/trend/",
                                   self._week_params(granularity="day"))
        self.assertEqual(response.status_code, 400)
