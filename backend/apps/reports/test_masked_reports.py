"""
脱敏报表测试

覆盖：用途按角色裁剪、字段裁剪、小样本合并/抑制、重叠窗口一致性、
迟到更正不改变已发布结果、跨权限用户结果一致、冻结地平线、窗口对齐、
导出审计与摘要、参数校验。
"""
import io
import json
from datetime import time, timedelta
from decimal import Decimal
from unittest.mock import patch

from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.authentication.backends import generate_token
from apps.authentication.models import User
from apps.warehouse.models import Category, Goods, StockIn, StockOut, Unit, Variety
from .masking import MASKING_POLICY, align_window
from .models import FrozenBucket, MaskedExportRecord, ReleasedCell


class MaskedReportTestBase(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user('masked-admin', 'testpass123', role='admin')
        self.external = User.objects.create_user('masked-external', 'testpass123', role='user')

        self.admin_client = APIClient()
        self.admin_client.credentials(HTTP_AUTHORIZATION=f'Bearer {generate_token(self.admin)}')
        self.external_client = APIClient()
        self.external_client.credentials(HTTP_AUTHORIZATION=f'Bearer {generate_token(self.external)}')

        unit = Unit.objects.create(name='件', created_by=self.admin)
        self.category_a = Category.objects.create(name='监管器材', unit=unit, created_by=self.admin)
        self.category_b = Category.objects.create(name='涉案物资', unit=unit, created_by=self.admin)
        self.variety_a1 = Variety.objects.create(name='封存设备', category=self.category_a, created_by=self.admin)
        self.variety_a2 = Variety.objects.create(name='扣押器材', category=self.category_a, created_by=self.admin)
        self.variety_b1 = Variety.objects.create(name='证物箱', category=self.category_b, created_by=self.admin)
        self.goods_a = Goods.objects.create(
            variety=self.variety_a1, name='封存终端甲', code='MSK-A-001',
            quantity=Decimal('100'), warning_threshold=Decimal('5'),
        )
        self.goods_b = Goods.objects.create(
            variety=self.variety_b1, name='证物箱乙', code='MSK-B-001',
            quantity=Decimal('100'), warning_threshold=Decimal('5'),
        )
        self.goods_c = Goods.objects.create(
            variety=self.variety_a2, name='扣押器材丙', code='MSK-C-001',
            quantity=Decimal('100'), warning_threshold=Decimal('5'),
        )
        self.today = timezone.localdate()

    # ------------------------------------------------------------------
    # 数据构造
    # ------------------------------------------------------------------

    def add_stock_in(self, goods, day, qty=1, times=1):
        for _ in range(times):
            record = StockIn.objects.create(
                goods=goods, operator=self.admin, quantity=qty,
                batch_no='批次-机密-1', supplier='某供应商',
            )
            StockIn.objects.filter(pk=record.pk).update(
                stock_in_time=timezone.make_aware(timezone.datetime.combine(day, time(12, 0)))
            )

    def add_stock_out(self, goods, day, qty=1, times=1):
        for _ in range(times):
            StockOut.objects.create(
                goods=goods, operator=self.admin, receiver='领用人甲',
                receiver_dept='某敏感部门', quantity=qty, status='completed',
                stock_out_time=timezone.make_aware(timezone.datetime.combine(day, time(13, 0))),
            )

    def get_trend(self, client, **params):
        return client.get('/api/masked-reports/trend/', params)

    def internal_params(self, start, end, **extra):
        params = {
            'purpose': 'internal_review',
            'start_date': start.isoformat(),
            'end_date': end.isoformat(),
            'granularity': 'day',
            'dimension': 'variety',
        }
        params.update(extra)
        return params

    def external_params(self, start, end, **extra):
        params = {
            'purpose': 'external_inspection',
            'start_date': start.isoformat(),
            'end_date': end.isoformat(),
        }
        params.update(extra)
        return params


class PurposeAndFieldPruningTest(MaskedReportTestBase):
    def test_purpose_list_filters_by_role(self):
        response = self.external_client.get('/api/masked-reports/purposes/')
        self.assertEqual(response.status_code, 200)
        purposes = [p['purpose'] for p in response.json()['data']]
        self.assertEqual(purposes, ['external_inspection'])

        response = self.admin_client.get('/api/masked-reports/purposes/')
        purposes = [p['purpose'] for p in response.json()['data']]
        self.assertIn('external_inspection', purposes)
        self.assertIn('internal_review', purposes)

    def test_external_report_prunes_detail_and_identity_fields(self):
        day = self.today - timedelta(days=10)
        self.add_stock_in(self.goods_a, day, times=4)
        self.add_stock_out(self.goods_a, day, times=1)

        response = self.get_trend(
            self.external_client,
            **self.external_params(self.today - timedelta(days=21), self.today),
        )
        self.assertEqual(response.status_code, 200)
        payload = json.dumps(response.json()['data'], ensure_ascii=False)

        # 明细与身份字段不得出现在报文中
        for leaked in ['领用人甲', '某敏感部门', '某供应商', '批次-机密-1',
                       '封存终端甲', 'MSK-A-001', 'masked-admin']:
            self.assertNotIn(leaked, payload)

        rows = response.json()['data']['rows']
        self.assertTrue(rows)
        allowed_keys = {'bucket', 'dimension_value', 'row_kind', 'in_count', 'in_quantity',
                        'out_count', 'out_quantity', 'record_count'}
        for row in rows:
            self.assertLessEqual(set(row.keys()), allowed_keys)
            # 外部用途只到品类粒度，不出现品种名
            self.assertNotIn(row['dimension_value'], ['封存设备', '扣押器材', '证物箱'])

    def test_internal_report_has_extra_field_and_variety_dimension(self):
        day = self.today - timedelta(days=5)
        self.add_stock_in(self.goods_a, day, times=5)

        response = self.get_trend(self.admin_client, **self.internal_params(day, day))
        self.assertEqual(response.status_code, 200)
        rows = response.json()['data']['rows']
        normal = [r for r in rows if r['row_kind'] == 'normal']
        self.assertEqual(len(normal), 1)
        self.assertEqual(normal[0]['dimension_value'], '封存设备')
        self.assertEqual(normal[0]['goods_type_count'], 1)
        self.assertEqual(normal[0]['record_count'], 5)


class SmallCellHandlingTest(MaskedReportTestBase):
    def test_small_groups_suppressed_below_threshold(self):
        day = self.today - timedelta(days=5)
        self.add_stock_in(self.goods_a, day, times=5)   # 品种A1：5 条，正常发布
        self.add_stock_in(self.goods_b, day, times=1)   # 品种B1：1 条，小组
        self.add_stock_in(self.goods_c, day, times=1)   # 品种A2：1 条，小组（合并后仍不足 k）

        response = self.get_trend(self.admin_client, **self.internal_params(day, day))
        self.assertEqual(response.status_code, 200)
        data = response.json()['data']
        rows = [r for r in data['rows'] if r['bucket'] == day.isoformat()]

        normal = [r for r in rows if r['row_kind'] == 'normal']
        self.assertEqual(len(normal), 1)
        self.assertEqual(normal[0]['dimension_value'], '封存设备')
        self.assertEqual(normal[0]['record_count'], 5)

        # 合并后仍不足 k=3：只出现不具名抑制标记，无任何数值
        suppressed = [r for r in rows if r['row_kind'] == 'suppressed']
        self.assertEqual(len(suppressed), 1)
        self.assertEqual(suppressed[0]['dimension_value'], '已抑制（小组）')
        for field in ['in_count', 'in_quantity', 'out_count', 'out_quantity', 'record_count']:
            self.assertIsNone(suppressed[0][field])

        # 所有发布行都满足 k 阈值；合计不含被抑制记录
        for row in rows:
            if row['row_kind'] != 'suppressed':
                self.assertGreaterEqual(row['record_count'], MASKING_POLICY['k_threshold'])
        self.assertEqual(data['totals']['grand']['record_count'], 5)
        self.assertEqual(data['summary']['suppressed_cell_count'], 2)

    def test_merged_group_released_when_reaching_threshold(self):
        day = self.today - timedelta(days=5)
        self.add_stock_in(self.goods_b, day, times=1)
        self.add_stock_in(self.goods_c, day, times=2)

        response = self.get_trend(self.admin_client, **self.internal_params(day, day))
        rows = response.json()['data']['rows']
        merged = [r for r in rows if r['row_kind'] == 'merged']
        self.assertEqual(len(merged), 1)
        self.assertEqual(merged[0]['dimension_value'], '其他（已合并）')
        self.assertEqual(merged[0]['record_count'], 3)
        self.assertEqual(response.json()['data']['totals']['grand']['record_count'], 3)

    def test_suppress_strategy_marks_small_cells_without_values(self):
        day = self.today - timedelta(days=5)
        self.add_stock_in(self.goods_a, day, times=5)
        self.add_stock_in(self.goods_b, day, times=1)

        with patch.dict(MASKING_POLICY['purposes']['internal_review'], {'strategy': 'suppress'}):
            response = self.get_trend(self.admin_client, **self.internal_params(day, day))
        self.assertEqual(response.status_code, 200)
        data = response.json()['data']
        rows = data['rows']

        suppressed = [r for r in rows if r['row_kind'] == 'suppressed']
        self.assertEqual(len(suppressed), 1)
        # 品种维度（仅内部）下标记可具名，但绝不发布数值
        self.assertEqual(suppressed[0]['dimension_value'], '证物箱')
        self.assertIsNone(suppressed[0]['record_count'])
        self.assertEqual(data['totals']['grand']['record_count'], 5)

    def test_cross_dimension_totals_exclude_suppressed_records(self):
        day = self.today - timedelta(days=5)
        self.add_stock_in(self.goods_a, day, times=3)   # 品种A1 达标
        self.add_stock_in(self.goods_c, day, times=1)   # 品种A2 被抑制

        variety_resp = self.get_trend(self.admin_client, **self.internal_params(day, day, dimension='variety'))
        category_resp = self.get_trend(self.admin_client, **self.internal_params(day, day, dimension='category'))
        self.assertEqual(variety_resp.status_code, 200)
        self.assertEqual(category_resp.status_code, 200)

        variety_total = variety_resp.json()['data']['totals']['grand']['record_count']
        category_total = category_resp.json()['data']['totals']['grand']['record_count']
        # 两个维度的合计一致，且都不包含被抑制的那 1 条记录
        self.assertEqual(variety_total, 3)
        self.assertEqual(category_total, 3)

        category_rows = category_resp.json()['data']['rows']
        normal = [r for r in category_rows if r['row_kind'] == 'normal']
        self.assertEqual(len(normal), 1)
        self.assertEqual(normal[0]['dimension_value'], '监管器材')
        self.assertEqual(normal[0]['record_count'], 3)


class ConsistencyTest(MaskedReportTestBase):
    def test_overlapping_windows_yield_identical_shared_buckets(self):
        for offset in range(6, 11):
            self.add_stock_in(self.goods_a, self.today - timedelta(days=offset), times=3)

        resp1 = self.get_trend(
            self.admin_client,
            **self.internal_params(self.today - timedelta(days=10), self.today - timedelta(days=6)),
        )
        resp2 = self.get_trend(
            self.admin_client,
            **self.internal_params(self.today - timedelta(days=8), self.today - timedelta(days=4)),
        )
        self.assertEqual(resp1.status_code, 200)
        self.assertEqual(resp2.status_code, 200)

        rows1 = {r['bucket']: r for r in resp1.json()['data']['rows']}
        rows2 = {r['bucket']: r for r in resp2.json()['data']['rows']}
        shared = set(rows1) & set(rows2)
        self.assertTrue(shared)
        for bucket in shared:
            self.assertEqual(rows1[bucket], rows2[bucket])

    def test_late_correction_does_not_change_frozen_results(self):
        day = self.today - timedelta(days=6)
        self.add_stock_in(self.goods_a, day, times=3)

        first = self.get_trend(self.admin_client, **self.internal_params(day, day))
        self.assertEqual(first.status_code, 200)
        first_rows = first.json()['data']['rows']
        first_fingerprint = first.json()['data']['summary']['data_fingerprint']
        self.assertEqual(FrozenBucket.objects.get(bucket_date=day).record_count, 3)

        # 迟到更正：同日又补录入库与出库
        self.add_stock_in(self.goods_a, day, times=2)
        self.add_stock_out(self.goods_a, day, times=1)

        second = self.get_trend(self.admin_client, **self.internal_params(day, day))
        self.assertEqual(second.status_code, 200)
        self.assertEqual(second.json()['data']['rows'], first_rows)
        self.assertEqual(second.json()['data']['summary']['data_fingerprint'], first_fingerprint)
        # 冻结快照不被更正改变
        self.assertEqual(FrozenBucket.objects.get(bucket_date=day).record_count, 3)
        cells = ReleasedCell.objects.filter(bucket_date=day)
        self.assertEqual(sum(c.in_count for c in cells), 3)

    def test_same_report_identical_across_roles(self):
        day = self.today - timedelta(days=10)
        self.add_stock_in(self.goods_a, day, times=4)

        params = self.external_params(self.today - timedelta(days=21), self.today)
        external_resp = self.get_trend(self.external_client, **params)
        admin_resp = self.get_trend(self.admin_client, **params)
        self.assertEqual(external_resp.status_code, 200)
        self.assertEqual(admin_resp.status_code, 200)

        external_data = external_resp.json()['data']
        admin_data = admin_resp.json()['data']
        self.assertEqual(external_data['rows'], admin_data['rows'])
        self.assertEqual(external_data['totals'], admin_data['totals'])
        self.assertEqual(
            external_data['summary']['data_fingerprint'],
            admin_data['summary']['data_fingerprint'],
        )

    def test_repeated_requests_are_stable(self):
        day = self.today - timedelta(days=5)
        self.add_stock_in(self.goods_a, day, times=4)

        first = self.get_trend(self.admin_client, **self.internal_params(day, day))
        second = self.get_trend(self.admin_client, **self.internal_params(day, day))
        self.assertEqual(first.json()['data']['rows'], second.json()['data']['rows'])
        self.assertEqual(
            first.json()['data']['summary']['data_fingerprint'],
            second.json()['data']['summary']['data_fingerprint'],
        )


class WindowAlignmentTest(MaskedReportTestBase):
    def test_horizon_clamps_recent_days(self):
        response = self.get_trend(
            self.admin_client,
            **self.internal_params(self.today - timedelta(days=3), self.today),
        )
        self.assertEqual(response.status_code, 200)
        summary = response.json()['data']['summary']
        horizon = self.today - timedelta(days=MASKING_POLICY['embargo_days'])
        self.assertEqual(summary['effective_window']['end'], horizon.isoformat())
        self.assertTrue(any('冻结规则' in note for note in summary['window_notes']))

    def test_window_without_releasable_data_returns_400(self):
        response = self.get_trend(
            self.admin_client,
            **self.internal_params(self.today, self.today),
        )
        self.assertEqual(response.status_code, 400)

    def test_week_alignment_shrinks_to_full_weeks(self):
        base = self.today - timedelta(days=20)
        monday = base + timedelta(days=(7 - base.weekday()) % 7)
        start = monday + timedelta(days=1)  # 周二，强制不对齐
        end = self.today - timedelta(days=1)

        response = self.get_trend(self.external_client, **self.external_params(start, end))
        self.assertEqual(response.status_code, 200)
        summary = response.json()['data']['summary']
        effective_start = summary['effective_window']['start']
        effective_end = summary['effective_window']['end']
        self.assertEqual(effective_start, (monday + timedelta(days=7)).isoformat())
        self.assertEqual(
            timezone.datetime.strptime(effective_end, '%Y-%m-%d').date().weekday(), 6
        )
        self.assertTrue(any('对齐' in note for note in summary['window_notes']))

    def test_align_window_never_expands_beyond_request(self):
        base = self.today - timedelta(days=30)
        for granularity in ['day', 'week', 'month']:
            aligned_start, aligned_end = align_window(base, self.today, granularity)
            self.assertGreaterEqual(aligned_start, base)
            self.assertLessEqual(aligned_end, self.today)

    def test_month_alignment_uses_full_months(self):
        # 请求窗口跨不满整月时收缩到完整月
        start = (self.today.replace(day=1) - timedelta(days=45)).replace(day=15)
        end = self.today - timedelta(days=1)
        response = self.get_trend(
            self.external_client,
            **self.external_params(start, end, granularity='month'),
        )
        self.assertEqual(response.status_code, 200)
        summary = response.json()['data']['summary']
        effective_start = timezone.datetime.strptime(summary['effective_window']['start'], '%Y-%m-%d').date()
        effective_end = timezone.datetime.strptime(summary['effective_window']['end'], '%Y-%m-%d').date()
        self.assertEqual(effective_start.day, 1)
        self.assertEqual((effective_end + timedelta(days=1)).day, 1)
        for row in response.json()['data']['rows']:
            self.assertRegex(row['bucket'], r'^\d{4}-\d{2}$')

    def test_empty_window_returns_empty_rows(self):
        day = self.today - timedelta(days=5)
        response = self.get_trend(self.admin_client, **self.internal_params(day, day))
        self.assertEqual(response.status_code, 200)
        data = response.json()['data']
        self.assertEqual(data['rows'], [])
        self.assertEqual(data['totals']['grand']['record_count'], 0)
        self.assertEqual(data['summary']['bucket_count'], 1)


class ExportAuditTest(MaskedReportTestBase):
    def test_preview_and_export_create_audit_records(self):
        day = self.today - timedelta(days=10)
        self.add_stock_in(self.goods_a, day, times=4)
        params = self.external_params(self.today - timedelta(days=21), self.today)

        preview = self.get_trend(self.external_client, **params)
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(MaskedExportRecord.objects.filter(action='preview').count(), 1)

        export = self.external_client.get('/api/masked-reports/export/', params)
        self.assertEqual(export.status_code, 200)
        record = MaskedExportRecord.objects.get(action='export')
        self.assertEqual(record.requested_by, self.external)
        self.assertEqual(record.requester_role, 'user')
        self.assertEqual(record.policy_version, MASKING_POLICY['version'])
        self.assertEqual(record.k_threshold, MASKING_POLICY['k_threshold'])
        self.assertEqual(len(record.data_fingerprint), 64)
        self.assertEqual(record.data_fingerprint, preview.json()['data']['summary']['data_fingerprint'])
        self.assertIn('领用人', record.fields_withheld)
        self.assertEqual(record.summary['effective_window']['start'], record.effective_start.isoformat())

    def test_export_xlsx_contains_summary_and_data_sheets(self):
        day = self.today - timedelta(days=10)
        self.add_stock_in(self.goods_a, day, times=4)

        response = self.external_client.get(
            '/api/masked-reports/export/',
            self.external_params(self.today - timedelta(days=21), self.today),
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response['Content-Type'])

        wb = load_workbook(io.BytesIO(response.content))
        self.assertEqual(wb.sheetnames, ['导出摘要', '收发趋势'])

        summary_sheet = wb['导出摘要']
        summary_text = '\n'.join(
            str(cell.value) for row in summary_sheet.iter_rows() for cell in row if cell.value is not None
        )
        self.assertIn('策略版本', summary_text)
        self.assertIn(MASKING_POLICY['version'], summary_text)
        self.assertIn('数据指纹', summary_text)
        self.assertIn('裁剪字段', summary_text)
        self.assertIn('实际数据范围', summary_text)

        data_sheet = wb['收发趋势']
        headers = [cell.value for cell in data_sheet[1]]
        self.assertEqual(headers, ['时间桶', '品类', '行类型', '入库次数', '入库数量', '出库次数', '出库数量', '记录数'])
        # 数据页不得出现明细与身份字段
        sheet_text = '\n'.join(
            str(cell.value) for row in data_sheet.iter_rows() for cell in row if cell.value is not None
        )
        for leaked in ['领用人甲', '某供应商', '封存终端甲', 'MSK-A-001']:
            self.assertNotIn(leaked, sheet_text)

    def test_records_endpoint_admin_only(self):
        day = self.today - timedelta(days=10)
        self.add_stock_in(self.goods_a, day, times=4)
        self.get_trend(self.external_client, **self.external_params(self.today - timedelta(days=21), self.today))

        forbidden = self.external_client.get('/api/masked-reports/records/')
        self.assertEqual(forbidden.status_code, 403)

        allowed = self.admin_client.get('/api/masked-reports/records/')
        self.assertEqual(allowed.status_code, 200)
        records = allowed.json()['data']
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]['policy_version'], MASKING_POLICY['version'])
        self.assertEqual(records[0]['requested_by'], 'masked-external')
        self.assertIn('summary', records[0])


class ParamValidationTest(MaskedReportTestBase):
    def test_invalid_params(self):
        day = self.today - timedelta(days=5)

        # 缺少用途
        response = self.admin_client.get('/api/masked-reports/trend/')
        self.assertEqual(response.status_code, 400)

        # 无效用途
        response = self.get_trend(self.admin_client, **self.internal_params(day, day, purpose='nope'))
        self.assertEqual(response.status_code, 400)

        # 角色无权使用用途
        response = self.get_trend(self.external_client, **self.internal_params(day, day))
        self.assertEqual(response.status_code, 403)

        # 外部用途不支持日粒度
        response = self.get_trend(
            self.external_client,
            **self.external_params(day, day, granularity='day'),
        )
        self.assertEqual(response.status_code, 400)

        # 外部用途不支持品种维度
        response = self.get_trend(
            self.external_client,
            **self.external_params(day, day, dimension='variety'),
        )
        self.assertEqual(response.status_code, 400)

        # 日期格式错误
        response = self.get_trend(
            self.admin_client,
            **self.internal_params(day, day, start_date='2026/01/01'),
        )
        self.assertEqual(response.status_code, 400)

        # 开始晚于结束
        response = self.get_trend(
            self.admin_client,
            **self.internal_params(day, day - timedelta(days=1)),
        )
        self.assertEqual(response.status_code, 400)
