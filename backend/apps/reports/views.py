"""
报表视图
"""
import logging
import os
import shutil
from datetime import datetime, timedelta
from django.db.models import Sum, Count, Q, F
from django.http import HttpResponse
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from apps.core.response import success_response, error_response
from apps.warehouse.models import Goods, StockIn, StockOut, Warning
from .anonymization import (
    AnonymizationError, build_trend_report, describe_policy,
)
from .models import DailyReport

logger = logging.getLogger('apps')


class DashboardView(APIView):
    """仪表盘数据"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        today = datetime.now().date()
        
        # 货物统计
        goods_count = Goods.objects.filter(is_active=True).count()
        warning_goods_count = Goods.objects.filter(
            is_active=True,
            quantity__lte=F('warning_threshold')
        ).count()
        
        # 今日入库统计
        today_in = StockIn.objects.filter(
            stock_in_time__date=today
        ).aggregate(
            count=Count('id'),
            total=Sum('quantity')
        )
        
        # 今日出库统计
        today_out = StockOut.objects.filter(
            stock_out_time__date=today,
            status='completed'
        ).aggregate(
            count=Count('id'),
            total=Sum('quantity')
        )
        
        # 待审批数量
        pending_approval_count = StockOut.objects.filter(status='pending').count()
        
        # 未读预警数量
        unread_warning_count = Warning.objects.filter(is_read=False).count()
        
        # 最近7天入库趋势
        in_trend = []
        out_trend = []
        for i in range(6, -1, -1):
            date = today - timedelta(days=i)
            in_data = StockIn.objects.filter(
                stock_in_time__date=date
            ).aggregate(total=Sum('quantity'))
            out_data = StockOut.objects.filter(
                stock_out_time__date=date,
                status='completed'
            ).aggregate(total=Sum('quantity'))
            
            in_trend.append({
                'date': date.strftime('%m-%d'),
                'value': float(in_data['total'] or 0)
            })
            out_trend.append({
                'date': date.strftime('%m-%d'),
                'value': float(out_data['total'] or 0)
            })
        
        # 库存预警货物列表
        warning_goods = Goods.objects.filter(
            is_active=True,
            quantity__lte=F('warning_threshold')
        ).values('id', 'name', 'code', 'quantity', 'warning_threshold')[:5]
        
        return success_response(data={
            'goods_count': goods_count,
            'warning_goods_count': warning_goods_count,
            'today_in_count': today_in['count'] or 0,
            'today_in_total': float(today_in['total'] or 0),
            'today_out_count': today_out['count'] or 0,
            'today_out_total': float(today_out['total'] or 0),
            'pending_approval_count': pending_approval_count,
            'unread_warning_count': unread_warning_count,
            'in_trend': in_trend,
            'out_trend': out_trend,
            'warning_goods': list(warning_goods)
        })


class DailyReportView(APIView):
    """每日报表"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        start_date = request.query_params.get('start_date')
        end_date = request.query_params.get('end_date')
        
        if not start_date:
            start_date = (datetime.now() - timedelta(days=30)).date()
        else:
            start_date = datetime.strptime(start_date, '%Y-%m-%d').date()
        
        if not end_date:
            end_date = datetime.now().date()
        else:
            end_date = datetime.strptime(end_date, '%Y-%m-%d').date()
        
        reports = []
        current_date = start_date
        
        while current_date <= end_date:
            # 入库统计
            in_data = StockIn.objects.filter(
                stock_in_time__date=current_date
            ).aggregate(
                count=Count('id'),
                total=Sum('quantity')
            )
            
            # 出库统计
            out_data = StockOut.objects.filter(
                stock_out_time__date=current_date,
                status='completed'
            ).aggregate(
                count=Count('id'),
                total=Sum('quantity')
            )
            
            # 预警统计
            warning_count = Warning.objects.filter(
                created_at__date=current_date
            ).count()
            
            reports.append({
                'date': current_date.strftime('%Y-%m-%d'),
                'in_count': in_data['count'] or 0,
                'in_total': float(in_data['total'] or 0),
                'out_count': out_data['count'] or 0,
                'out_total': float(out_data['total'] or 0),
                'warning_count': warning_count
            })
            
            current_date += timedelta(days=1)
        
        return success_response(data=reports)


class ExportView(APIView):
    """数据导出"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        export_type = request.query_params.get('type', 'goods')
        
        wb = Workbook()
        ws = wb.active
        
        # 样式定义
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='0066FF', end_color='0066FF', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        if export_type == 'goods':
            ws.title = '货物清单'
            headers = ['编码', '名称', '品类', '品种', '规格', '库存', '单位', '存放位置', '预警阈值']
            ws.append(headers)
            
            # 设置表头样式
            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border
            
            # 数据
            goods_list = Goods.objects.select_related('variety', 'variety__category', 'unit').filter(is_active=True)
            for goods in goods_list:
                ws.append([
                    goods.code,
                    goods.name,
                    goods.variety.category.name if goods.variety else '',
                    goods.variety.name if goods.variety else '',
                    goods.specification,
                    float(goods.quantity),
                    goods.unit.name if goods.unit else '',
                    goods.location,
                    float(goods.warning_threshold)
                ])
            
            filename = f'货物清单_{datetime.now().strftime("%Y%m%d%H%M%S")}.xlsx'
            
        elif export_type == 'stock_in':
            ws.title = '入库记录'
            headers = ['货物编码', '货物名称', '入库数量', '单位', '批次号', '供应商', '操作人', '入库时间']
            ws.append(headers)
            
            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border
            
            stock_ins = StockIn.objects.select_related('goods', 'goods__unit', 'operator').all()
            for record in stock_ins:
                ws.append([
                    record.goods.code,
                    record.goods.name,
                    float(record.quantity),
                    record.goods.unit.name if record.goods.unit else '',
                    record.batch_no,
                    record.supplier,
                    record.operator.username if record.operator else '',
                    record.stock_in_time.strftime('%Y-%m-%d %H:%M:%S')
                ])
            
            filename = f'入库记录_{datetime.now().strftime("%Y%m%d%H%M%S")}.xlsx'
            
        elif export_type == 'stock_out':
            ws.title = '出库记录'
            headers = ['货物编码', '货物名称', '出库数量', '单位', '领用人', '领用部门', '状态', '操作人', '出库时间']
            ws.append(headers)
            
            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border
            
            stock_outs = StockOut.objects.select_related('goods', 'goods__unit', 'operator').all()
            for record in stock_outs:
                ws.append([
                    record.goods.code,
                    record.goods.name,
                    float(record.quantity),
                    record.goods.unit.name if record.goods.unit else '',
                    record.receiver,
                    record.receiver_dept,
                    record.get_status_display(),
                    record.operator.username if record.operator else '',
                    record.stock_out_time.strftime('%Y-%m-%d %H:%M:%S') if record.stock_out_time else ''
                ])
            
            filename = f'出库记录_{datetime.now().strftime("%Y%m%d%H%M%S")}.xlsx'
        else:
            return error_response(message='无效的导出类型')
        
        # 调整列宽
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2) * 1.2
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # 返回文件
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        wb.save(response)
        
        logger.info(f"Data exported: {export_type}")
        return response


class SystemMonitorView(APIView):
    """返回容器内可直接读取的轻量运行状态。"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        cpu_count = os.cpu_count() or 1
        load_1m = os.getloadavg()[0] if hasattr(os, "getloadavg") else 0.0
        disk_total, disk_used, _ = shutil.disk_usage('/')
        disk_percent = round(disk_used * 100 / disk_total, 2) if disk_total else 0
        process_count = len([name for name in os.listdir('/proc') if name.isdigit()]) if os.path.isdir('/proc') else 1
        
        return success_response(data={
            'cpu': {
                'percent': round(load_1m * 100 / cpu_count, 2),
                'count': cpu_count
            },
            'memory': {
                'total': None,
                'used': None,
                'percent': None
            },
            'disk': {
                'total': round(disk_total / (1024 ** 3), 2),
                'used': round(disk_used / (1024 ** 3), 2),
                'percent': disk_percent
            },
            'network': {
                'bytes_sent': None,
                'bytes_recv': None
            },
            'system': {
                'boot_time': None,
                'uptime': None,
                'process_count': process_count
            }
        })


# ==================== 脱敏趋势报表 ====================

def _parse_date(value, name):
    try:
        return datetime.strptime(value, '%Y-%m-%d').date()
    except (TypeError, ValueError):
        raise AnonymizationError(f'{name} 格式应为 YYYY-MM-DD')


def _log_anonymized_export(request, report):
    """记录脱敏报表导出操作，含策略版本与数据范围，供审计追溯。"""
    from apps.authentication.models import OperationLog
    summary = report['summary']
    try:
        OperationLog.objects.create(
            user=request.user,
            action='export',
            module='脱敏报表',
            detail=(
                f"用途={summary['purpose']} 策略版本={summary['strategy_version']} "
                f"窗口={summary['aligned_window']['start']}~{summary['aligned_window']['end']} "
                f"快照={summary['snapshot_ids']}"
            ),
            ip_address=request.META.get('REMOTE_ADDR'),
            user_agent=request.headers.get('User-Agent', '')[:500]
        )
    except Exception as e:
        logger.error(f"Failed to log anonymized export: {e}")


class AnonymizedTrendReportView(APIView):
    """
    按用途脱敏的收发趋势报表。

    GET 参数：
    - purpose: 报表用途（external_inspection / internal_audit）
    - granularity: 统计粒度（week / month）
    - start_date / end_date: 请求窗口，向外对齐到整周期，进行中周期不输出
    - export: xlsx 时导出文件（附导出摘要供审核人确认），默认返回 JSON
    """
    permission_classes = [IsAuthenticated]

    def get(self, request):
        params = request.query_params
        purpose = params.get('purpose', 'external_inspection')
        granularity = params.get('granularity', 'week')

        try:
            end_date = _parse_date(params.get('end_date'), 'end_date') \
                if params.get('end_date') else datetime.now().date()
            start_date = _parse_date(params.get('start_date'), 'start_date') \
                if params.get('start_date') else end_date - timedelta(days=28)
            report = build_trend_report(
                request.user, purpose, granularity, start_date, end_date
            )
        except AnonymizationError as e:
            return error_response(message=e.message, code=e.code)

        _log_anonymized_export(request, report)

        if params.get('export') == 'xlsx':
            return self._render_xlsx(report)
        return success_response(data=report)

    def _render_xlsx(self, report):
        """导出 xlsx：第一个工作表为导出摘要，第二个为脱敏数据。"""
        summary = report['summary']
        wb = Workbook()

        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='0066FF', end_color='0066FF', fill_type='solid')

        # 摘要工作表：数据范围与处理规则，供审核人确认
        ws_summary = wb.active
        ws_summary.title = '导出摘要'
        ws_summary.append(['项目', '内容'])
        for col in range(1, 3):
            cell = ws_summary.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
        strategy = summary['strategy']
        summary_rows = [
            ('报表用途', f"{summary['purpose_label']}（{summary['purpose']}）"),
            ('脱敏策略版本', summary['strategy_version']),
            ('请求窗口', f"{summary['requested_window']['start']} ~ {summary['requested_window']['end']}"),
            ('对齐后数据范围', f"{summary['aligned_window']['start']} ~ {summary['aligned_window']['end']}"),
            ('统计粒度', strategy['granularity']),
            ('进行中周期已排除', '是' if summary['current_period_excluded'] else '否'),
            ('小样本阈值(组合并/抑制)', strategy['min_group_size']),
            ('小样本处理方式', strategy['suppression_mode']),
            ('总量取整步长', strategy['rounding_step']),
            ('数据口径', summary['record_scope']),
            ('输出字段', '、'.join(summary['fields'])),
            ('已裁剪敏感字段', '、'.join(summary['excluded_sensitive_fields'])),
            ('合并组数 / 抑制组数', f"{summary['totals']['merged_groups']} / {summary['totals']['suppressed_groups']}"),
            ('快照编号', ', '.join(str(i) for i in summary['snapshot_ids'])),
            ('生成时间', summary['generated_at']),
        ]
        for row in summary_rows:
            ws_summary.append(list(row))
        ws_summary.column_dimensions['A'].width = 24
        ws_summary.column_dimensions['B'].width = 80

        # 数据工作表
        ws_data = wb.create_sheet('趋势数据')
        headers = ['周期'] + [
            {'category': '品类', 'variety': '品种'}[d]
            for d in summary['dimensions'] if d != 'period'
        ] + ['入库次数', '入库总量', '出库次数', '出库总量']
        ws_data.append(headers)
        for col in range(1, len(headers) + 1):
            cell = ws_data.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
        dim_keys = [d for d in summary['dimensions'] if d != 'period']
        for row in report['rows']:
            ws_data.append(
                [row['period']]
                + [row.get(d, '') for d in dim_keys]
                + [row['in_count'], row['in_total'], row['out_count'], row['out_total']]
            )
        for column in ws_data.columns:
            ws_data.column_dimensions[column[0].column_letter].width = 18

        filename = f"脱敏趋势报表_{summary['purpose']}_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        wb.save(response)
        logger.info(f"Anonymized report exported: {summary['purpose']} {summary['strategy_version']}")
        return response


class AnonymizationPolicyView(APIView):
    """返回当前脱敏策略与用途授权清单，供审核人确认处理规则。"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return success_response(data=describe_policy(request.user))
